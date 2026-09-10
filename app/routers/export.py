from fastapi import APIRouter, Depends
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os
from app.database import get_db
from app.models import Patient, Visits, Prescription, PrescriptionItem, Bill, Medicine, User
from app.core.auth import get_current_user
import tempfile
import os 
router = APIRouter()

@router.get("all")
def export_all_data(
    db:Session = Depends(get_db),
    current_user:User = Depends(get_current_user)
):

        # Create workbook
    wb = openpyxl.Workbook()

    header_font = Font(bold=True , color="FFFFFF")
    header_fill = PatternFill(
        start_color="1F4E76",
        end_color="1F4E79",
        fill_type="solid"
    )

    center = Alignment(horizontal="center")


    def style_header(ws,headers):
        for col,header in enumerate(headers,1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        # ─── SHEET 1: PATIENTS ───

    ws1 = wb.active
    ws1.title = "Patients"

    patient_headers = [
        "Patient Id","Name","Phone",
        "Gender","Date of Birth","Address",
        "Registered Date"
    ]
    style_header(ws1,patient_headers)

    patients = db.query(Patient).filter(
        Patient.clinic_id == current_user.clinic_id
    ).all()

    for row,p in enumerate(patients,2):
        ws1.cell(row=row , column=1,value=p.patient_id)
        ws1.cell(row=row, column=2,value=p.name)
        ws1.cell(row=row,column=3,value=p.phone)
        ws1.cell(row=row,column=4,value=p.gender)
        ws1.cell(row=row, column=5, value=str(p.date_of_birth) if p.date_of_birth else "")
        ws1.cell(row=row,column=6,value=p.address or "")
        ws1.cell(row=row,column=7,value=str(p.created_at.date()) if p.created_at else "")


        # Auto column width
        for col in ws1.columns:

            max_len = max(len(str(cell.value or "")) for cell in col)
            ws1.column_dimensions[col[0].column_letter].width = max_len + 4


    # ─── SHEET 2: VISITS ───
    ws2 = wb.create_sheet("Visits")

    visit_headers = [
        "Visit ID", "Patient ID", "Patient Name",
        "Doctor ID", "Complaint", "Diagnosis",
        "Notes", "Follow-up Date", "Date"
    ]

    style_header(ws2, visit_headers)

    visits = db.query(Visits).filter(
        Visits.clinic_id == current_user.clinic_id
    ).all()

    for row, v in enumerate(visits, 2):
        patient = db.query(Patient).filter(
            Patient.patient_id == v.patient_id
        ).first()

        ws2.cell(row=row, column=1, value=v.id)
        ws2.cell(row=row, column=2, value=v.patient_id)
        ws2.cell(row=row, column=3, value=patient.name if patient else "Unknown")
        ws2.cell(row=row, column=4, value=v.doctor_id)
        ws2.cell(row=row, column=5, value=v.complaint or "")
        ws2.cell(row=row, column=6, value=v.diagnosis or "")
        ws2.cell(row=row, column=7, value=v.notes or "")
        ws2.cell(row=row, column=8, value=str(v.follow_up_date) if v.follow_up_date else "")
        ws2.cell(row=row, column=9, value=str(v.created_at.date()) if v.created_at else "")

        for col in ws2.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws2.column_dimensions[col[0].column_letter].width = max_len + 4

          # ─── SHEET 3: PRESCRIPTIONS ───
    ws3 = wb.create_sheet("Prescriptions")

    prescription_headers = [
        "Prescription ID", "Patient ID", "Patient Name",
        "Doctor ID", "Medicine", "Frequency",
        "Duration", "Quantity", "Timing", "Status", "Date"
    ]
    style_header(ws3, prescription_headers)
    prescriptions = db.query(Prescription).filter(
        Prescription.clinic_id == current_user.clinic_id
    ).all()

    row = 2
    for p in prescriptions:
        patient = db.query(Patient).filter(
            Patient.patient_id == p.patient_id
        ).first()

        items = db.query(PrescriptionItem).filter(
            PrescriptionItem.prescription_id == p.id
        ).all()

        for item in items:
            medicine = db.query(Medicine).filter(
                Medicine.id == item.medicine_id
            ).first()

            ws3.cell(row=row, column=1, value=p.id)
            ws3.cell(row=row, column=2, value=p.patient_id)
            ws3.cell(row=row, column=3, value=patient.name if patient else "Unknown")
            ws3.cell(row=row, column=4, value=p.doctor_id)
            ws3.cell(row=row, column=5, value=medicine.name if medicine else "Unknown")
            ws3.cell(row=row, column=6, value=item.frequency)
            ws3.cell(row=row, column=7, value=item.duration)
            ws3.cell(row=row, column=8, value=item.quantity)
            ws3.cell(row=row, column=9, value=item.timing or "")
            ws3.cell(row=row, column=10, value=p.status)
            ws3.cell(row=row, column=11, value=str(p.created_at.date()) if p.created_at else "")
            row += 1

    for col in ws3.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws3.column_dimensions[col[0].column_letter].width = max_len + 4

    # ─── SHEET 4: BILLS ───
    ws4 = wb.create_sheet("Bills")

    bill_headers = [
        "Bill ID", "Patient ID", "Patient Name",
        "Registration Fee", "Consultation Fee",
        "Medicine Total", "Discount",
        "Total Amount", "Payment Method",
        "Status", "Date"
    ]
    style_header(ws4, bill_headers)

    bills = db.query(Bill).filter(
        Bill.clinic_id == current_user.clinic_id
    ).all()

    for row, b in enumerate(bills, 2):
        patient = db.query(Patient).filter(
            Patient.patient_id == b.patient_id
        ).first()

        ws4.cell(row=row, column=1, value=b.id)
        ws4.cell(row=row, column=2, value=b.patient_id)
        ws4.cell(row=row, column=3, value=patient.name if patient else "Unknown")
        ws4.cell(row=row, column=4, value=b.registration_fee)
        ws4.cell(row=row, column=5, value=b.consultation_fee)
        ws4.cell(row=row, column=6, value=b.medicine_total)
        ws4.cell(row=row, column=7, value=b.discount)
        ws4.cell(row=row, column=8, value=b.total_amount)
        ws4.cell(row=row, column=9, value=b.payment_method or "")
        ws4.cell(row=row, column=10, value=b.status)
        ws4.cell(row=row, column=11, value=str(b.created_at.date()) if b.created_at else "")

    for col in ws4.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws4.column_dimensions[col[0].column_letter].width = max_len + 4

    # ─── SAVE FILE ───
    filename = f"HMS_AI_Export_{date.today()}.xlsx"
    filepath = os.path.join(tempfile.gettempdir(), filename)
    wb.save(filepath)

    return FileResponse(
        path=filepath,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )






       
